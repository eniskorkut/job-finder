"""CV upload and text extraction: real signatures, previews, ownership."""

from __future__ import annotations

import io

import pymupdf
import pytest
from docx import Document

from app.integrations.cv_extraction import ExtractionStatus, extract, sniff_kind


def make_pdf(text: str | None = "Enis Korkut - AI Engineer\nPython, PyTorch, FastAPI, LLM, RAG") -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 96), text, fontsize=11)
    payload = document.tobytes()
    document.close()
    return payload


def make_docx(text: str = "Enis Korkut - AI Engineer\nPython, PyTorch, FastAPI, LLM") -> bytes:
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Deneyim"
    table.rows[0].cells[1].text = "6 yıl"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def upload(api, name: str, content: bytes, content_type: str):
    return api.post(
        "/api/v1/cvs", files={"file": (name, io.BytesIO(content), content_type)}
    )


class TestExtractionUnit:
    def test_sniff_detects_real_formats(self):
        assert sniff_kind(make_pdf()) == "pdf"
        assert sniff_kind(make_docx()) == "docx"
        assert sniff_kind(b"plain text") == "text"
        assert sniff_kind(b"MZ\x90\x00binary") == "binary"

    def test_pdf_with_text(self):
        result = extract(filename="cv.pdf", content=make_pdf())
        assert result.status == ExtractionStatus.OK
        assert "PyTorch" in (result.text or "")
        assert result.pages == 1

    def test_scanned_pdf_asks_for_ocr(self):
        result = extract(filename="tarama.pdf", content=make_pdf(text=None))
        assert result.status == ExtractionStatus.OCR_REQUIRED
        assert result.text is None
        assert "OCR" in (result.warning or "")

    def test_docx_extracts_paragraphs_and_tables(self):
        result = extract(filename="cv.docx", content=make_docx())
        assert result.status == ExtractionStatus.OK
        assert "PyTorch" in (result.text or "")
        assert "6 yıl" in (result.text or "")

    def test_legacy_doc_is_unsupported(self):
        ole_header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
        result = extract(filename="eski.doc", content=ole_header)
        assert result.status == ExtractionStatus.UNSUPPORTED
        assert ".doc" in (result.warning or "")

    def test_fake_extension_is_detected_by_content(self):
        result = extract(filename="cv.docx", content=b"%PDF-1.4 not really a docx")
        # Content wins over the extension: it is parsed as a PDF (and fails cleanly).
        assert result.status in {ExtractionStatus.FAILED, ExtractionStatus.OCR_REQUIRED}

    def test_short_text_is_flagged_empty(self):
        result = extract(filename="bos.txt", content="kısa".encode())
        assert result.status == ExtractionStatus.EMPTY


class TestUploadAndPreview:
    def test_pdf_upload_extracts_and_preview_shows_text(self, api_user1):
        response = upload(api_user1, "cv.pdf", make_pdf(), "application/pdf")
        assert response.status_code == 201, response.text
        body = response.json()
        assert body["extraction_status"] == ExtractionStatus.OK.value
        assert body["has_extracted_text"] is True
        assert body["extraction_warning"] is None

        preview = api_user1.get(f"/api/v1/cvs/{body['id']}/preview")
        assert preview.status_code == 200
        payload = preview.json()
        assert "PyTorch" in payload["text"]
        assert payload["character_count"] > 20
        assert payload["line_count"] >= 1

    def test_docx_upload_extracts(self, api_user1):
        response = upload(
            api_user1,
            "cv.docx",
            make_docx(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        assert response.status_code == 201
        assert response.json()["extraction_status"] == ExtractionStatus.OK.value

    def test_scanned_pdf_upload_warns_instead_of_faking_text(self, api_user1):
        response = upload(api_user1, "taranmis.pdf", make_pdf(text=None), "application/pdf")
        assert response.status_code == 201
        body = response.json()
        assert body["extraction_status"] == ExtractionStatus.OCR_REQUIRED.value
        assert body["has_extracted_text"] is False
        assert "OCR" in body["extraction_warning"]

        preview = api_user1.get(f"/api/v1/cvs/{body['id']}/preview").json()
        assert preview["text"] == ""
        assert preview["extraction_status"] == ExtractionStatus.OCR_REQUIRED.value

    def test_fake_extension_is_rejected(self, api_user1):
        """A PDF renamed to .docx (with the docx MIME type) must not pass."""
        response = upload(
            api_user1,
            "cv.docx",
            make_pdf(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        assert response.status_code == 422
        assert "eşleşmiyor" in response.json()["detail"]["message"]

    def test_short_pdf_text_is_flagged_not_ocr(self, api_user1):
        response = upload(api_user1, "kisa.pdf", make_pdf("Kısa CV"), "application/pdf")
        assert response.status_code == 201
        assert response.json()["extraction_status"] == ExtractionStatus.EMPTY.value

    def test_legacy_doc_is_rejected_with_a_clear_message(self, api_user1):
        response = upload(
            api_user1, "cv.doc", b"\xd0\xcf\x11\xe0" + b"\x00" * 128, "application/msword"
        )
        assert response.status_code == 422
        assert "Legacy" in response.json()["detail"]["message"]

    def test_size_limit(self, api_user1):
        oversized = b"a" * (10 * 1024 * 1024 + 1)
        response = upload(api_user1, "buyuk.txt", oversized, "text/plain")
        assert response.status_code == 422
        assert "10 MB" in response.json()["detail"]["message"]

    def test_preview_is_owner_only(self, api_user1, api_user2):
        created = upload(api_user1, "gizli.pdf", make_pdf("Gizli icerik LLM Engineer"), "application/pdf")
        cv_id = created.json()["id"]

        assert api_user2.get(f"/api/v1/cvs/{cv_id}/preview").status_code == 404
        assert api_user2.post(f"/api/v1/cvs/{cv_id}/extract").status_code == 404
        assert api_user2.get(f"/api/v1/cvs/{cv_id}").status_code == 404
        assert api_user2.get(f"/api/v1/cvs/{cv_id}/download").status_code == 404

    def test_extracted_text_is_never_returned_in_the_listing(self, api_user1):
        upload(api_user1, "liste.pdf", make_pdf("Gizli icerik: LLM Engineer, Python, PyTorch"), "application/pdf")
        listing = api_user1.get("/api/v1/cvs").json()
        assert listing[0]["has_extracted_text"] is True
        assert "text" not in listing[0]

    def test_reextract_endpoint(self, api_user1):
        created = upload(api_user1, "yeniden.pdf", make_pdf(), "application/pdf").json()
        response = api_user1.post(f"/api/v1/cvs/{created['id']}/extract")
        assert response.status_code == 200
        assert "PyTorch" in response.json()["text"]

    def test_second_upload_deactivates_the_first(self, api_user1):
        first = upload(api_user1, "ilk.pdf", make_pdf(), "application/pdf").json()
        second = upload(api_user1, "ikinci.pdf", make_pdf(), "application/pdf").json()
        listing = {cv["id"]: cv for cv in api_user1.get("/api/v1/cvs").json()}
        assert listing[first["id"]]["is_active"] is False
        assert listing[second["id"]]["is_active"] is True


@pytest.mark.parametrize("filename", ["cv.pdf", "cv.docx"])
def test_extraction_does_not_touch_the_network(api_user1, monkeypatch, filename):
    """CV text extraction is local; DeepSeek profiling is phase 3."""
    calls: list[str] = []

    def explode(*args, **kwargs):  # pragma: no cover - must never run
        calls.append("called")
        raise AssertionError("Harici bir çağrı yapılmamalı.")

    import httpx

    monkeypatch.setattr(httpx.AsyncClient, "request", explode)
    content = make_pdf() if filename.endswith("pdf") else make_docx()
    response = upload(api_user1, filename, content, "application/pdf" if filename.endswith("pdf") else "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.status_code == 201
    assert calls == []
