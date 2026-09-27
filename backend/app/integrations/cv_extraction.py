"""CV text extraction.

Only formats whose real signature we can verify are parsed: the file
extension and the client supplied MIME type are never trusted. A scanned PDF
is reported as ``ocr_required`` instead of silently producing an empty CV.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

from enum import StrEnum

MIN_TEXT_CHARS = 40


class ExtractionStatus(StrEnum):
    OK = "ok"
    OCR_REQUIRED = "ocr_required"
    UNSUPPORTED = "unsupported"
    EMPTY = "empty"
    FAILED = "failed"


@dataclass(slots=True)
class ExtractionResult:
    text: str | None
    status: ExtractionStatus
    warning: str | None = None
    pages: int | None = None


def sniff_kind(content: bytes) -> str:
    """Detect the real container format from its magic bytes."""
    if content.startswith(b"%PDF-"):
        return "pdf"
    if content.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile:
            return "unknown"
        if "word/document.xml" in names:
            return "docx"
        if "[Content_Types].xml" in names:
            return "ooxml-other"
        return "zip"
    if content.startswith(b"{\\rtf"):
        return "rtf"
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"  # legacy .doc / .xls container
    if b"\x00" in content[:2048]:
        return "binary"
    return "text"


def extract(
    *, filename: str, content: bytes, content_type: str | None = None
) -> ExtractionResult:
    kind = sniff_kind(content)
    suffix = Path(filename).suffix.lower()

    if kind == "pdf":
        return _extract_pdf(content)
    if kind == "docx":
        return _extract_docx(content)
    if kind in {"ole", "ooxml-other", "rtf", "zip", "unknown", "binary"}:
        return ExtractionResult(
            text=None,
            status=ExtractionStatus.UNSUPPORTED,
            warning=(
                "Bu dosya biçimi metin çıkarımı için desteklenmiyor. "
                "PDF, DOCX, TXT veya MD yükleyin. (Legacy .doc desteklenmez.)"
            ),
        )

    # text/plain and friends: extension must still be a text one
    if suffix not in {".txt", ".md", ""}:
        return ExtractionResult(
            text=None,
            status=ExtractionStatus.UNSUPPORTED,
            warning="Dosya içeriği desteklenen bir biçimle eşleşmiyor.",
        )
    return _decode_text(content)


def _decode_text(content: bytes) -> ExtractionResult:
    for encoding in ("utf-8", "cp1254", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 never fails
        return ExtractionResult(None, ExtractionStatus.FAILED, "Metin çözümlenemedi.")

    cleaned = text.strip()
    if len(cleaned) < MIN_TEXT_CHARS:
        return ExtractionResult(
            text=cleaned or None,
            status=ExtractionStatus.EMPTY,
            warning="CV metni çok kısa; içerik boş görünüyor.",
        )
    return ExtractionResult(text=cleaned, status=ExtractionStatus.OK)


def _extract_pdf(content: bytes) -> ExtractionResult:
    try:
        import pymupdf  # imported lazily so the API can boot without it
    except ImportError:  # pragma: no cover - dependency is declared
        return ExtractionResult(
            None, ExtractionStatus.FAILED, "PDF ayrıştırıcı yüklenemedi."
        )

    try:
        document = pymupdf.open(stream=content, filetype="pdf")
    except Exception:
        return ExtractionResult(
            None, ExtractionStatus.FAILED, "PDF dosyası açılamadı (bozuk olabilir)."
        )

    try:
        pages = document.page_count
        chunks = [page.get_text("text") or "" for page in document]
    finally:
        document.close()

    text = "\n".join(chunks).strip()
    if len(text) < 5:
        # No text layer at all: almost certainly a scan.
        return ExtractionResult(
            text=None,
            status=ExtractionStatus.OCR_REQUIRED,
            warning=(
                "PDF metin içermiyor (taranmış görüntü olabilir). "
                "OCR bu sürümde desteklenmiyor; metin tabanlı bir PDF yükleyin."
            ),
            pages=pages,
        )
    if len(text) < MIN_TEXT_CHARS:
        return ExtractionResult(
            text=text,
            status=ExtractionStatus.EMPTY,
            warning="PDF metni çok kısa; CV içeriği eksik görünüyor.",
            pages=pages,
        )
    return ExtractionResult(text=text, status=ExtractionStatus.OK, pages=pages)


def _extract_docx(content: bytes) -> ExtractionResult:
    try:
        import docx  # python-docx
    except ImportError:  # pragma: no cover - dependency is declared
        return ExtractionResult(
            None, ExtractionStatus.FAILED, "DOCX ayrıştırıcı yüklenemedi."
        )

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception:
        return ExtractionResult(
            None, ExtractionStatus.FAILED, "DOCX dosyası açılamadı (bozuk olabilir)."
        )

    parts: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(part.strip() for part in parts if part and part.strip()).strip()
    if len(text) < MIN_TEXT_CHARS:
        return ExtractionResult(
            text=None,
            status=ExtractionStatus.EMPTY,
            warning="DOCX içinde okunabilir metin bulunamadı.",
        )
    return ExtractionResult(text=text, status=ExtractionStatus.OK)
